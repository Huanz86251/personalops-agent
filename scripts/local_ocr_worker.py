"""One request per process: document-scoped lazy OCR, no server or Agent imports.

Invoked with -I -S. Only the dedicated dependency directory is added to sys.path.
The parent enforces the deadline and terminates/reaps this process on timeout.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import os
import sys
import warnings
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEPS = ROOT / ".agent" / "ocr-deps"
MAX_PIXELS = 12_000_000
MAX_OCR_IMAGES = 60
MAX_RESULT_CHARS = 2_000_000


def detect_format(data: bytes, name: str) -> str:
    """Content signatures select a parser; the parser must still validate input."""
    import filetype
    from PIL import Image

    if b"%PDF-" in data[:1024]:
        return "pdf"
    if data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 10000
                or sum(e.file_size for e in entries) > 100 * 1024**2
            ):
                raise ValueError("Expanded archive exceeds local limits.")
            names = set(archive.namelist())
            for marker, kind in (
                ("ppt/presentation.xml", "pptx"),
                ("word/document.xml", "docx"),
                ("xl/workbook.xml", "xlsx"),
            ):
                if marker in names:
                    return kind
        raise ValueError(
            "Unsupported archive; ZIP files are not automatically unpacked for OCR."
        )
    guessed = filetype.guess(data)
    if guessed and guessed.mime.startswith("image/"):
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        return "image"
    # Text formats have no reliable universal signature; validate bytes as UTF-8.
    suffix = Path(name).suffix.lower().lstrip(".")
    if suffix in {"txt", "md", "html", "csv", "json"}:
        text = data.decode("utf-8-sig")
        if "\x00" not in text:
            return suffix
    raise ValueError(
        "Unknown or unsupported file content; OCR requires a decodable image or document."
    )


def load_image(data: bytes):
    from PIL import Image, ImageOps

    Image.MAX_IMAGE_PIXELS = MAX_PIXELS
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(data)) as source:
            if source.width * source.height > MAX_PIXELS:
                raise ValueError(
                    "Image exceeds 12 million pixels; crop or downsize it first."
                )
            if getattr(source, "n_frames", 1) != 1:
                raise ValueError(
                    "Multi-frame images are not supported; export the intended frame first."
                )
            return ImageOps.exif_transpose(source).convert("RGB")


class Reading:
    """Shares one lazy engine and image cache for the entire document invocation."""

    def __init__(self, mode="auto", engine_factory=None):
        self.mode = mode
        self.engine = None
        self.engine_error = None
        self.engine_factory = engine_factory or self._new_engine
        self.engine_loads = 0
        self.ocr_calls = 0
        self.cache = {}
        self.blocks = []
        self.warnings = []
        self.total_chars = 0

    @staticmethod
    def _new_engine():
        import cv2
        import rapidocr
        from rapidocr import (
            EngineType,
            LangDet,
            LangRec,
            ModelType,
            OCRVersion,
            RapidOCR,
        )

        cv2.setNumThreads(2)
        model_dir = Path(rapidocr.__file__).parent / "models"

        def model(pattern):
            matches = list(model_dir.glob(pattern))
            if len(matches) != 1:
                raise ValueError(
                    "Local OCR weights missing; run scripts/setup_ocr.py before reading."
                )
            return str(matches[0])

        # Explicit local paths prevent inference-time weight downloads.
        return RapidOCR(
            params={
                "Det.engine_type": EngineType.ONNXRUNTIME,
                "Det.lang_type": LangDet.CH,
                "Det.model_type": ModelType.SMALL,
                "Det.ocr_version": OCRVersion.PPOCRV6,
                "Det.model_path": model("*v6_det_small*.onnx"),
                "Rec.engine_type": EngineType.ONNXRUNTIME,
                "Rec.lang_type": LangRec.CH,
                "Rec.model_type": ModelType.SMALL,
                "Rec.ocr_version": OCRVersion.PPOCRV6,
                "Rec.model_path": model("*v6_rec_small*.onnx"),
                "Cls.model_path": model("*cls*.onnx"),
                "EngineConfig.onnxruntime.intra_op_num_threads": 2,
                "EngineConfig.onnxruntime.inter_op_num_threads": 1,
                "EngineConfig.onnxruntime.use_cuda": False,
                "Global.text_score": 0.0,
                "Global.log_level": "error",
            }
        )

    def warn(self, location, message):
        self.warnings.append(f"{location}: {message}")

    def native(self, location, text):
        if text.strip():
            self.total_chars += len(text)
            if self.total_chars > MAX_RESULT_CHARS:
                raise ValueError(
                    "Reading output too large; request a smaller page range."
                )
            self.blocks.append({"source": "native", "location": location, "text": text})

    def ocr(self, location, image, region=None):
        if self.mode == "off":
            self.warn(location, "OCR disabled; image text has not been read.")
            return
        try:
            if isinstance(image, bytes):
                image = load_image(image)
            if image.width * image.height > MAX_PIXELS:
                raise ValueError("Image exceeds OCR pixel limit.")
            key = hashlib.sha256(image.tobytes() + str(image.size).encode()).hexdigest()
            if key not in self.cache:
                if self.ocr_calls >= MAX_OCR_IMAGES:
                    raise ValueError(
                        "OCR image budget reached; read a smaller page range."
                    )
                if self.engine is None:
                    if self.engine_error is not None:
                        raise ValueError(self.engine_error)
                    try:
                        self.engine = self.engine_factory()
                        self.engine_loads += 1
                    except Exception as error:
                        self.engine_error = (
                            f"OCR engine unavailable: {str(error)[:250]}"
                        )
                        raise
                import numpy as np

                self.ocr_calls += 1
                # RapidOCR ndarray input is OpenCV BGR, whereas Pillow is RGB.
                output = self.engine(
                    np.asarray(image.convert("RGB"))[:, :, ::-1].copy()
                )
                texts = output.txts or []
                scores = output.scores if output.scores is not None else []
                boxes = output.boxes if output.boxes is not None else []
                self.cache[key] = [
                    {
                        "text": str(text),
                        "score": round(float(score), 4),
                        "box": [[round(float(v), 2) for v in point] for point in box],
                    }
                    for text, score, box in zip(texts, scores, boxes)
                ]
            lines = self.cache[key]
            text = "\n".join(line["text"] for line in lines)
            self.total_chars += len(text)
            if self.total_chars > MAX_RESULT_CHARS:
                raise ValueError("OCR output budget reached.")
            self.blocks.append(
                {
                    "source": "ocr",
                    "location": location,
                    "region": region,
                    "text": text,
                    "lines": lines,
                    "pixel_size": [image.width, image.height],
                }
            )
            if not lines:
                self.warn(
                    location,
                    "No text recognized; this does not mean the image is empty or understood.",
                )
            elif any(line["score"] < 0.8 for line in lines):
                self.warn(
                    location,
                    "Some OCR scores are below 0.8; verify uncertain characters/numbers.",
                )
        except Exception as error:  # noqa: BLE001 -- keep partial reading and surface engine errors
            self.warn(
                location, f"OCR failed: {type(error).__name__}: {str(error)[:300]}"
            )

    def result(self, kind, total_units=1, start=1, count=1):
        end = min(total_units, start + count - 1)
        next_page = end + 1 if end < total_units else None
        if start > 1 or next_page:
            self.warn(
                "Document",
                f"Only units {start}-{end} of {total_units} were requested/read.",
            )
        limitations = [
            "OCR extracts text only: scenes, chart trends and diagram relationships are not understood.",
            "OCR scores are engine scores, not guaranteed accuracy; table cell structure is not guaranteed.",
        ]
        summary = {
            "format": kind,
            "status": "partial" if self.warnings else "read",
            "worker_pid": os.getpid(),
            "total_units": total_units,
            "start_page": start,
            "end_page": end,
            "next_page": next_page,
            "ocr_images": self.ocr_calls,
            "engine_loads": self.engine_loads,
            "warnings": self.warnings,
            "limitations": limitations,
        }
        sections = [
            "# Attachment reading",
            json.dumps(summary, ensure_ascii=False, indent=2),
        ]
        for block in self.blocks:
            label = "OCR 提取文字" if block["source"] == "ocr" else "文档原生文字"
            sections.append(
                f"## {block['location']} — {label}\n\n{block['text'] or '[未识别到文字]'}"
            )
        return {
            "summary": summary,
            "blocks": self.blocks,
            "markdown": "\n\n".join(sections),
        }


def page_range(total, start, count):
    if not 1 <= start <= total:
        raise ValueError(f"start_page must be between 1 and {total}.")
    return range(start - 1, min(total, start - 1 + count))


def read_pdf(data, reader, start, count):
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        total = len(pdf.pages)
        for index in page_range(total, start, count):
            page = pdf.pages[index]
            location = f"Page {index + 1}"
            try:
                text = page.extract_text(layout=False) or ""
                reader.native(location, text)
                # Render only pages requiring OCR. Rendering also covers outlines,
                # masks and PDF composition that raw image extraction can miss.
                suspect = len(text.strip()) < 8 or "\ufffd" in text or "(cid:" in text
                full = reader.mode == "force" or (suspect and bool(page.objects))
                regions = []
                if full:
                    regions = [(location + " full-page", page.bbox)]
                else:
                    for number, obj in enumerate(page.images, 1):
                        x0, y0, x1, y1 = page.bbox
                        box = (
                            max(x0, obj["x0"]),
                            max(y0, obj["top"]),
                            min(x1, obj["x1"]),
                            min(y1, obj["bottom"]),
                        )
                        if box[2] > box[0] and box[3] > box[1]:
                            regions.append((f"{location}, image {number}", box))
                    if page.curves and not page.images:
                        reader.warn(
                            location,
                            "Vector graphics detected; their visual meaning has not been read.",
                        )
                if regions and reader.mode != "off":
                    # Limit allocation before rasterizing unusually large pages.
                    pixels = max(1.0, page.width * page.height)
                    dpi = min(200, math.sqrt(MAX_PIXELS / pixels) * 72 * 0.98)
                    rendered = page.to_image(resolution=dpi).original.convert("RGB")
                    sx, sy = rendered.width / page.width, rendered.height / page.height
                    for region_location, box in regions[: MAX_OCR_IMAGES + 1]:
                        x0, top, _, _ = page.bbox
                        crop = rendered.crop(
                            (
                                int((box[0] - x0) * sx),
                                int((box[1] - top) * sy),
                                int((box[2] - x0) * sx),
                                int((box[3] - top) * sy),
                            )
                        )
                        reader.ocr(region_location, crop, list(box))
                    if len(regions) > MAX_OCR_IMAGES + 1:
                        reader.warn(
                            location, "Remaining image regions skipped due to budget."
                        )
                elif regions:
                    reader.warn(location, "OCR disabled; image text has not been read.")
            except Exception as error:  # noqa: BLE001 -- report the failed page, continue the document
                reader.warn(
                    location,
                    f"Page reading failed: {type(error).__name__}: {str(error)[:300]}",
                )
            finally:
                page.close()
    return reader.result("pdf", total, start, count)


def read_pptx(data, reader, start, count):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))
    total = len(presentation.slides)

    def visit(shapes, location, picture_parts):
        for index, shape in enumerate(shapes, 1):
            loc = f"{location}, shape {index}"
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, loc, picture_parts)
                continue
            if shape.has_text_frame:
                reader.native(loc, shape.text)
            if shape.has_table:
                reader.native(
                    loc,
                    "\n".join(
                        " | ".join(c.text for c in row.cells)
                        for row in shape.table.rows
                    ),
                )
            if shape.has_chart:
                reader.warn(
                    loc,
                    "Native chart is not a raster picture; chart data/visual interpretation not extracted.",
                )
            # Picture fills/backgrounds are reported below via relationships.
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    picture_parts.add(hashlib.sha256(shape.image.blob).hexdigest())
                    image = load_image(shape.image.blob)
                    width, height = image.size
                    image = image.crop(
                        (
                            int(width * shape.crop_left),
                            int(height * shape.crop_top),
                            int(width * (1 - shape.crop_right)),
                            int(height * (1 - shape.crop_bottom)),
                        )
                    )
                    reader.ocr(loc, image)
                except Exception as error:  # noqa: BLE001 -- unsupported/corrupt images remain visible as warnings
                    reader.warn(loc, f"Picture unreadable: {str(error)[:250]}")

    for index in page_range(total, start, count):
        slide = presentation.slides[index]
        picture_parts = set()
        visit(slide.shapes, f"Slide {index + 1}", picture_parts)
        # Covers embedded picture fills/backgrounds as well as normal pictures.
        # The image-content cache avoids repeated model inference.
        for rel in slide.part.rels.values():
            if rel.reltype.endswith("/image"):
                if rel.is_external:
                    reader.warn(
                        f"Slide {index + 1}", "External image not downloaded/read."
                    )
                elif (
                    hashlib.sha256(rel.target_part.blob).hexdigest()
                    not in picture_parts
                ):
                    reader.ocr(
                        f"Slide {index + 1}, embedded {rel.rId}", rel.target_part.blob
                    )
        if slide.has_notes_slide:
            reader.native(
                f"Slide {index + 1}, notes", slide.notes_slide.notes_text_frame.text
            )
    reader.warn(
        "PPTX",
        "Master/layout artwork, SmartArt and vector drawings are not rendered; image OCR is text only.",
    )
    return reader.result("pptx", total, start, count)


def read_other(data, kind, reader):
    from markitdown import StreamInfo, converters

    converter_name = {
        "docx": "DocxConverter",
        "xlsx": "XlsxConverter",
        "html": "HtmlConverter",
        "csv": "CsvConverter",
    }.get(kind, "PlainTextConverter")
    converter = getattr(converters, converter_name)()
    reader.native(
        "Document",
        converter.convert(
            io.BytesIO(data), StreamInfo(extension="." + kind)
        ).text_content,
    )
    if kind in {"docx", "xlsx"}:
        prefix = "word/media/" if kind == "docx" else "xl/media/"
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            media = [
                n
                for n in archive.namelist()
                if n.startswith(prefix) and not n.endswith("/")
            ]
            for name in media[:MAX_OCR_IMAGES]:
                reader.ocr(f"Embedded image {name}", archive.read(name))
            if len(media) > MAX_OCR_IMAGES:
                reader.warn(
                    "Document", "Remaining embedded images skipped due to budget."
                )
            if media:
                reader.warn(
                    "Document",
                    "Image OCR is an appendix by package path; page/cell placement is not reconstructed.",
                )
            if any(
                "/charts/" in name or "/diagrams/" in name
                for name in archive.namelist()
            ):
                reader.warn(
                    "Document",
                    "Native charts/diagrams are not raster images; visual content is not rendered.",
                )
            # Never fetch linked images or silently suggest that they were read.
            from defusedxml import ElementTree

            for name in archive.namelist():
                if name.endswith(".rels"):
                    for rel in ElementTree.fromstring(archive.read(name)):
                        if (
                            rel.get("Type", "").endswith("/image")
                            and rel.get("TargetMode") == "External"
                        ):
                            reader.warn(name, "External image not downloaded/read.")
    if kind == "html":
        reader.warn(
            "HTML",
            "Linked/embedded HTML images are not read; save an image as a task file and use ocr_image.",
        )
    return reader.result(kind)


def run(request, engine_factory=None):
    data = Path(request["input_path"]).read_bytes()
    if len(data) > 20 * 1024**2:
        raise ValueError("Input exceeds 20 MiB.")
    kind = detect_format(data, request["name"])
    if request.get("operation") == "inspect":
        return {"summary": {"format": kind}, "blocks": [], "markdown": ""}
    reader = Reading(request.get("ocr", "auto"), engine_factory)
    if kind == "json":
        reader.native("Document", data.decode("utf-8-sig"))
        return reader.result(kind)
    if request.get("operation") == "image" and kind != "image":
        raise ValueError(
            "ocr_image accepts an image; use attachment_to_text for documents."
        )
    start, count = request.get("start_page", 1), request.get("max_pages", 20)
    if (
        not isinstance(start, int)
        or not isinstance(count, int)
        or start < 1
        or not 1 <= count <= 20
    ):
        raise ValueError("Invalid page range; max_pages must be 1..20.")
    if kind == "pdf":
        return read_pdf(data, reader, start, count)
    if kind == "pptx":
        return read_pptx(data, reader, start, count)
    if start != 1:
        raise ValueError("Page selection applies only to PDF/PPTX.")
    if kind == "image":
        image = load_image(data)
        crop = request.get("crop")
        if crop is not None:
            if (
                len(crop) != 4
                or not all(isinstance(v, int) for v in crop)
                or not (
                    0 <= crop[0] < crop[2] <= image.width
                    and 0 <= crop[1] < crop[3] <= image.height
                )
            ):
                raise ValueError(
                    "crop must be [left, top, right, bottom] pixels within the oriented image."
                )
            image = image.crop(crop)
        reader.ocr("Image", image, crop)
        return reader.result("image")
    return read_other(data, kind, reader)


def main():
    sys.path.insert(0, str(DEPS))
    # No document/image is uploaded and no missing model is fetched at runtime.
    import socket

    def no_network(*args, **kwargs):
        raise OSError("Network disabled in local attachment reader.")

    socket.socket.connect = no_network
    socket.create_connection = no_network
    request = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    try:
        result = run(request)
    except Exception as error:  # noqa: BLE001 -- process boundary returns structured failures to parent
        result = {"error": f"{type(error).__name__}: {str(error)[:800]}"}
    Path(sys.argv[2]).write_text(
        json.dumps(result, ensure_ascii=False), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
