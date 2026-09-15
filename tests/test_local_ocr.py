"""Real, provider-free OCR/document tests plus process/file boundary checks."""

import asyncio
import base64
import hashlib
import io
import json
import subprocess
import sys
import unittest
import zipfile
import zlib
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from deepagents.backends.utils import create_file_data
from langchain_core.utils.function_calling import convert_to_openai_tool

from scripts import local_ocr_worker as worker
from tools import local_native as local


def image_bytes(text="中文测试 销售金额 1200 元", english="English OCR Total 123.45"):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1000, 240), "white")
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 42)
    draw = ImageDraw.Draw(image)
    draw.text((30, 30), text, font=font, fill="black")
    draw.text((30, 115), english, font=font, fill="black")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def pdf_bytes(images):
    """Small real PDF with text and optional images; no PDF creation dependency."""
    from PIL import Image

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    page_ids = []
    for image_data in images:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(b"")
        content_id = len(objects) + 1
        image_id = content_id + 1
        content = b"BT /F1 18 Tf 35 740 Td (Native heading) Tj ET\n"
        resources = b"/Font << /F1 3 0 R >>"
        if image_data:
            content += b"q 500 0 0 120 35 540 cm /Im1 Do Q\n"
            resources += f" /XObject << /Im1 {image_id} 0 R >>".encode()
        objects.append(
            f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"endstream"
        )
        if image_data:
            image = Image.open(io.BytesIO(image_data)).convert("RGB")
            compressed = zlib.compress(image.tobytes())
            objects.append(
                (
                    f"<< /Type /XObject /Subtype /Image /Width {image.width} /Height {image.height} "
                    f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode /Length {len(compressed)} >>\nstream\n"
                ).encode()
                + compressed
                + b"\nendstream"
            )
        objects[page_id - 1] = (
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] "
                f"/Contents {content_id} 0 R /Resources << "
            ).encode()
            + resources
            + b" >> >>"
        )
    objects[1] = (
        f"<< /Type /Pages /Count {len(page_ids)} /Kids ["
        + " ".join(f"{p} 0 R" for p in page_ids)
        + "] >>"
    ).encode()
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return bytes(output)


class OCRContractTests(unittest.TestCase):
    def test_schema_and_worker_visibility(self):
        from tools import ALL_TOOLS
        from workers.general_worker import select_general_worker_tools
        from workers.web_worker import select_web_worker_tools

        for current in [local.attachment_to_text, local.ocr_image]:
            schema = convert_to_openai_tool(current)["function"]["parameters"]
            self.assertNotIn("runtime", schema["properties"])
            self.assertTrue(
                all(p.get("description") for p in schema["properties"].values())
            )
        for selected in [
            select_general_worker_tools(),
            select_web_worker_tools(ALL_TOOLS),
        ]:
            self.assertTrue(
                {"ocr_image", "attachment_to_text"}.issubset(t.name for t in selected)
            )
        with self.assertRaises(ValueError):
            local.attachment_to_text.args_schema.model_validate(
                {"path": "/a.pdf", "output_path": "/artifacts/a.md", "max_pages": 21}
            )

    def test_unknown_binary_and_non_image_rejected(self):
        with self.assertRaisesRegex(ValueError, "Unknown or unsupported"):
            local._reading_process(b"\x00\x01binary\xff", "file.bin", operation="image")
        with self.assertRaisesRegex(ValueError, "use attachment_to_text"):
            local._reading_process(pdf_bytes([None]), "file.bin", operation="image")

    def test_timeout_reports_release(self):
        with (
            patch.object(
                local.subprocess,
                "run",
                side_effect=subprocess.TimeoutExpired("reader", 1),
            ),
            self.assertRaisesRegex(ValueError, "timed out.*released"),
        ):
            local._reading_process(b"hello", "a.txt")

    def test_actual_timed_out_child_is_reaped(self):
        import psutil

        with TemporaryDirectory() as directory:
            script = Path(directory) / "slow_reader.py"
            script.write_text(
                "import os,time\nfrom pathlib import Path\nPath(__file__).with_suffix('.pid').write_text(str(os.getpid()))\ntime.sleep(30)\n",
                encoding="utf-8",
            )
            with (
                patch.object(local, "OCR_WORKER_SCRIPT", script),
                patch.object(local, "OCR_TIMEOUT_SECONDS", 1),
                self.assertRaisesRegex(ValueError, "timed out"),
            ):
                local._reading_process(b"hello", "a.txt")
            self.assertFalse(
                psutil.pid_exists(int(script.with_suffix(".pid").read_text()))
            )

    def test_download_is_owned_and_hash_checked(self):
        from tools.web_artifacts import _owner_directory

        data = b"owned"
        with TemporaryDirectory() as directory:
            root = Path(directory)
            target = _owner_directory(root, "web-a") / "file.bin"
            target.write_bytes(data)
            record = {
                "candidate_id": "download-1",
                "source_url": "https://example.com/file.bin",
                "filename": "file.bin",
                "tool_call_id": "download-call-1",
                "downloaded_at": "2026-09-06T00:00:00Z",
                "storage_path": str(target),
                "size_bytes": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
            runtime = SimpleNamespace(
                state={"worker_id": "web-a", "worker_downloaded_artifacts": [record]}
            )
            with patch("tools.web_artifacts.WEB_ARTIFACT_ROOT", root):
                self.assertEqual(local._read("/downloads/download-1", runtime), data)
                runtime.state["worker_id"] = "web-b"
                with self.assertRaises(ValueError):
                    local._read("/downloads/download-1", runtime)
                runtime.state["worker_id"] = "web-a"
                target.write_bytes(b"edited")
                with self.assertRaisesRegex(ValueError, "changed"):
                    local._read("/downloads/download-1", runtime)

    def test_engine_failure_preserves_native_text_with_warning(self):
        def fail():
            raise ValueError("missing model")

        reader = worker.Reading(engine_factory=fail)
        reader.native("page1", "Native text")
        reader.ocr("image1", image_bytes())
        result = reader.result("pdf")
        self.assertIn("Native text", result["markdown"])
        self.assertEqual(result["summary"]["status"], "partial")
        self.assertIn("missing model", result["summary"]["warnings"][0])


class OCRIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Test fixture dependencies only; production never imports OCR into Agent.
        sys.path.append(str(local.OCR_DEPENDENCY_ROOT))
        cls.image = image_bytes()

    def test_bilingual_image_tool_saves_json_and_releases_process(self):
        import psutil

        runtime = SimpleNamespace(
            state={
                "files": {
                    "/artifacts/table.bin": create_file_data(
                        base64.b64encode(self.image).decode(), encoding="base64"
                    )
                }
            },
            tool_call_id="ocr-test",
        )
        command = local.ocr_image.func(
            "/artifacts/table.bin", "/artifacts/read.md", runtime
        )
        message = json.loads(command.update["messages"][0].content)
        self.assertIn("中文测试", message["preview"])
        self.assertIn("English OCR Total 123.45", message["preview"])
        self.assertEqual(message["reading"]["engine_loads"], 1)
        self.assertTrue(message["reading"]["process_released"])
        self.assertFalse(psutil.pid_exists(message["reading"]["worker_pid"]))
        self.assertIn("/artifacts/read.json", command.update["files"])
        self.assertNotIn("onnxruntime", sys.modules)

    def test_pdf_mixed_pages_reuse_engine_and_native_page_needs_no_ocr(self):
        second = image_bytes("第二页 数量 42", "Second page 42")
        data = pdf_bytes([None, self.image, second])
        result = local._reading_process(data, "download.bin", operation="document")
        self.assertEqual(result["summary"]["format"], "pdf")
        self.assertEqual(result["summary"]["engine_loads"], 1)
        self.assertEqual(result["summary"]["ocr_images"], 2)
        self.assertIn("中文测试", result["markdown"])
        self.assertIn("第二页", result["markdown"])
        self.assertTrue(any(b["source"] == "native" for b in result["blocks"]))
        first = local._reading_process(data, "a.pdf", max_pages=1)
        self.assertEqual(first["summary"]["engine_loads"], 0)
        self.assertEqual(first["summary"]["next_page"], 2)

    def test_scanned_pdf_and_explicit_off(self):
        from PIL import Image

        stream = io.BytesIO()
        Image.open(io.BytesIO(self.image)).convert("RGB").save(stream, format="PDF")
        result = local._reading_process(stream.getvalue(), "scan.pdf")
        self.assertIn("中文测试", result["markdown"])
        self.assertEqual(result["blocks"][0]["source"], "ocr")
        off = local._reading_process(stream.getvalue(), "scan.pdf", ocr="off")
        self.assertEqual(off["summary"]["engine_loads"], 0)
        self.assertTrue(off["summary"]["warnings"])

    def test_pptx_text_and_two_pictures_reuse_engine(self):
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        for index in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(
                Inches(0.2), Inches(0.2), Inches(5), Inches(1)
            )
            box.text = f"Native slide {index + 1}"
            image = (
                self.image
                if index == 0
                else image_bytes("第二张图片 567", "Second image 567")
            )
            slide.shapes.add_picture(
                io.BytesIO(image), Inches(0.2), Inches(1.5), width=Inches(8)
            )
        stream = io.BytesIO()
        deck.save(stream)
        result = local._reading_process(stream.getvalue(), "file.bin")
        self.assertEqual(result["summary"]["engine_loads"], 1)
        self.assertEqual(result["summary"]["ocr_images"], 2)
        self.assertIn("Native slide 1", result["markdown"])
        self.assertIn("中文测试", result["markdown"])
        self.assertIn("第二张图片", result["markdown"])

    def test_repeated_images_cache_and_crop_bounds(self):
        result = local._reading_process(pdf_bytes([self.image, self.image]), "a.pdf")
        self.assertEqual(result["summary"]["ocr_images"], 1)
        self.assertEqual(len([b for b in result["blocks"] if b["source"] == "ocr"]), 2)
        with self.assertRaisesRegex(ValueError, "crop"):
            local._reading_process(
                self.image, "a.png", operation="image", crop=[0, 0, 10000, 10000]
            )

    def test_docx_and_xlsx_embedded_images(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr(
                "[Content_Types].xml",
                """<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>""",
            )
            archive.writestr(
                "_rels/.rels",
                """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>""",
            )
            archive.writestr(
                "word/document.xml",
                """<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Native document text</w:t></w:r></w:p></w:body></w:document>""",
            )
            archive.writestr("word/media/image1.png", self.image)
        result = local._reading_process(stream.getvalue(), "a.docx")
        self.assertIn("Native document text", result["markdown"])
        self.assertIn("中文测试", result["markdown"])
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image

        book = Workbook()
        book.active["A1"] = "Native cell"
        book.active.add_image(Image(io.BytesIO(self.image)), "B3")
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        result = local._reading_process(stream.getvalue(), "a.xlsx")
        self.assertIn("Native cell", result["markdown"])
        self.assertIn("中文测试", result["markdown"])

    def test_download_tool_to_real_ocr(self):
        import httpx

        from tools import web_artifacts

        with TemporaryDirectory() as directory:
            root = Path(directory)
            transport = httpx.MockTransport(
                lambda req: httpx.Response(200, content=self.image, request=req)
            )
            original = web_artifacts.download_to_record

            def download(**kwargs):
                return original(**kwargs, transport=transport)

            runtime = SimpleNamespace(
                state={"worker_id": "ocr-web", "files": {}},
                tool_call_id="download-ocr",
                stream_writer=lambda value: None,
            )
            tool = web_artifacts.create_download_web_artifact_tool(
                max_file_mib=1, root=root
            )
            with (
                patch.object(web_artifacts, "download_to_record", download),
                patch.object(web_artifacts, "WEB_ARTIFACT_ROOT", root),
            ):
                command = asyncio.run(
                    tool.coroutine("https://example.com/download.bin", runtime)
                )
                runtime.state["worker_downloaded_artifacts"] = command.update[
                    "worker_downloaded_artifacts"
                ]
                result = json.loads(command.update["messages"][0].content)
                read = local.ocr_image.func(
                    result["reading_path"], "/artifacts/download.md", runtime
                )
            self.assertIn(
                "中文测试", json.loads(read.update["messages"][0].content)["preview"]
            )


if __name__ == "__main__":
    unittest.main()
