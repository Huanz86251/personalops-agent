"""Offline document extraction inside the same restricted Code container."""

import argparse
import json
from pathlib import Path
import subprocess
import tempfile
import hashlib
import zipfile

MAX_BYTES = 20 * 1024 * 1024


def read(path, start_page=1, max_pages=20):
    source = Path(path).resolve()
    if not any(
        source.is_relative_to(Path(root))
        for root in ("/workspace", "/handoff", "/review")
    ):
        raise ValueError("Document must be in /workspace, /handoff or /review")
    if not source.is_file() or source.stat().st_size > MAX_BYTES:
        raise ValueError("Document missing or exceeds 20 MiB")
    if start_page < 1 or not 1 <= max_pages <= 20:
        raise ValueError("Invalid page range")
    warnings, sections = [], []
    suffix = source.suffix.lower()
    with source.open("rb") as stream:
        signature = stream.read(5)
    if suffix == ".pdf" or signature == b"%PDF-":
        import pymupdf as fitz

        with fitz.open(source) as document:
            if document.needs_pass:
                raise ValueError("Encrypted PDF requires a decrypted task copy")
            total = len(document)
            if start_page > total:
                raise ValueError("Page is outside document")
            end = min(total, start_page - 1 + max_pages)
            for index in range(start_page - 1, end):
                page = document[index]
                text = page.get_text().strip()
                # Scan and mixed-image pages need OCR, not an unsupported model attachment.
                if not text or page.get_images():
                    with tempfile.TemporaryDirectory(prefix="document-ocr-") as tmp:
                        image = Path(tmp) / "page.png"
                        if page.rect.width * page.rect.height * 2.25 > 12_000_000:
                            warnings.append(
                                f"Page {index + 1}: image too large for OCR"
                            )
                        else:
                            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                            pix.save(image)
                            result = subprocess.run(
                                [
                                    "tesseract",
                                    str(image),
                                    "stdout",
                                    "-l",
                                    "chi_sim+eng",
                                ],
                                capture_output=True,
                                timeout=60,
                            )
                            if result.returncode:
                                warnings.append(f"Page {index + 1}: OCR failed")
                            else:
                                text += "\n[local OCR]\n" + result.stdout.decode(
                                    "utf-8", errors="replace"
                                )
                sections.append(f"## Page {index + 1}\n{text}")
    elif suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(source) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 10000
                or sum(item.file_size for item in entries) > 100 * 1024 * 1024
            ):
                raise ValueError(
                    "Expanded Office archive exceeds 100 MiB / 10000 entries"
                )
        if start_page != 1:
            raise ValueError(
                "Office reader starts at 1; convert to PDF for page-specific reading"
            )
        total = end = 1
        if suffix == ".docx":
            from docx import Document

            document = Document(source)
            sections = [p.text for p in document.paragraphs]
            sections += [
                " | ".join(c.text for c in row.cells)
                for table in document.tables
                for row in table.rows
            ]
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            cell_chars = 0
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    sections.append("## Sheet " + sheet.title)
                    if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 100:
                        warnings.append(
                            f"Sheet {sheet.title}: limited to first 10000 rows and 100 columns"
                        )
                    for row in sheet.iter_rows(
                        max_row=min(sheet.max_row or 10000, 10000),
                        max_col=min(sheet.max_column or 100, 100),
                        values_only=True,
                    ):
                        row_text = " | ".join(
                            str(c) if c is not None else "" for c in row
                        )
                        cell_chars += len(row_text)
                        if cell_chars > MAX_BYTES:
                            raise ValueError("Extracted cells exceed local text limit")
                        sections.append(row_text)
            finally:
                workbook.close()
        else:
            from pptx import Presentation

            for index, slide in enumerate(Presentation(source).slides, 1):
                sections.append(f"## Slide {index}")
                sections += [
                    shape.text for shape in slide.shapes if shape.has_text_frame
                ]
        warnings.append(
            "Office embedded images/formula recalculation/layout not verified; use LibreOffice PDF export when needed"
        )
    else:
        raise ValueError("Unsupported document format")
    text = "\n\n".join(sections)
    if len(text.encode("utf-8")) > MAX_BYTES:
        raise ValueError("Extracted text exceeds 20 MiB; narrow the document first")
    truncated = len(text) > 24000
    cache = Path(tempfile.gettempdir()) / (
        "document-text-" + hashlib.sha256(text.encode("utf-8")).hexdigest() + ".md"
    )
    cache.write_text(text, encoding="utf-8")
    return {
        "path": path,
        "start_page": start_page,
        "end_page": end,
        "total_pages": total,
        "next_page": end + 1 if end < total else None,
        "warnings": warnings,
        "limitations": [
            "Local extraction/OCR reads text; it does not understand charts, scenes or diagram relationships"
        ],
        "truncated": truncated,
        "full_text_path": str(cache),
        "text": text[:24000],
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--start-page", type=int, default=1)
    parser.add_argument("--max-pages", type=int, default=20)
    args = parser.parse_args()
    try:
        print(
            json.dumps(
                read(args.path, args.start_page, args.max_pages), ensure_ascii=False
            )
        )
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        raise SystemExit(1)
